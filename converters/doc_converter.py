import io
import json
import zipfile
import xml.etree.ElementTree as ET
from typing import Dict, Any, Tuple
try:
    import pandas as pd
except ImportError:
    pd = None
from pypdf import PdfReader
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
import markdown
from bs4 import BeautifulSoup

DOC_MIME_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain",
    "md": "text/markdown",
    "html": "text/html",
    "csv": "text/csv",
    "tsv": "text/tab-separated-values",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "py": "text/x-python",
    "ipynb": "application/x-ipynb+json",
    "pages": "application/x-iwork-pages-sffpages",
    "numbers": "application/x-iwork-numbers-sffnumbers",
    "key": "application/x-iwork-keynote-sffkey",
    "webloc": "text/plain",
    "rtf": "application/rtf",
    "epub": "application/epub+zip"
}

SUPPORTED_DOC_FORMATS = list(DOC_MIME_TYPES.keys())


def convert_document(
    input_bytes: bytes,
    src_ext: str,
    target_ext: str,
    options: Dict[str, Any] = None
) -> Tuple[bytes, str, str]:
    """
    Handles Document, Spreadsheet, Jupyter Notebook, and Apple iWork conversions.
    Returns: (output_bytes, mime_type, target_ext)
    """
    if options is None:
        options = {}

    src_clean = src_ext.lower().strip().replace(".", "")
    target_clean = target_ext.lower().strip().replace(".", "")
    if target_clean == "markdown":
        target_clean = "md"

    if target_clean not in DOC_MIME_TYPES:
        target_clean = "txt"

    # 1. Apple iWork Bundles (.pages, .numbers, .key)
    if src_clean in ["pages", "numbers", "key"]:
        return _convert_iwork_bundle(input_bytes, src_clean, target_clean)

    # 2. Apple Safari Link (.webloc)
    if src_clean == "webloc":
        return _convert_webloc(input_bytes, target_clean)

    # 3. Jupyter Notebook (.ipynb) processing
    if src_clean == "ipynb":
        return _convert_ipynb(input_bytes, target_clean)

    # 4. Spreadsheet Processing (XLSX, XLS, CSV, TSV)
    if src_clean in ["xlsx", "xls", "csv", "tsv"] or target_clean in ["xlsx", "csv", "tsv"]:
        if src_clean in ["xlsx", "xls", "csv", "tsv"] and target_clean in ["xlsx", "csv", "tsv", "html", "txt", "md"]:
            return _convert_spreadsheet(input_bytes, src_clean, target_clean)

    # 5. PDF Input Processing
    if src_clean == "pdf":
        return _convert_from_pdf(input_bytes, target_clean)

    # 6. DOCX Input Processing
    if src_clean == "docx":
        return _convert_from_docx(input_bytes, target_clean)

    # 7. Text / Markdown / HTML / RTF Input Processing
    text_content = _extract_text_from_bytes(input_bytes, src_clean)

    if target_clean == "pdf":
        output_bytes = _render_text_to_pdf(text_content, src_clean)
        return output_bytes, DOC_MIME_TYPES["pdf"], "pdf"
    elif target_clean == "docx":
        output_bytes = _render_text_to_docx(text_content)
        return output_bytes, DOC_MIME_TYPES["docx"], "docx"
    elif target_clean == "html":
        if src_clean == "md":
            html_body = markdown.markdown(text_content)
            full_html = f"<!DOCTYPE html><html><head><meta charset='utf-8'></head><body>{html_body}</body></html>"
        else:
            escaped_text = text_content.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            full_html = f"<!DOCTYPE html><html><body><pre>{escaped_text}</pre></body></html>"
        return full_html.encode("utf-8"), DOC_MIME_TYPES["html"], "html"
    elif target_clean == "md":
        if src_clean == "html":
            soup = BeautifulSoup(text_content, "html.parser")
            md_text = soup.get_text()
        else:
            md_text = text_content
        return md_text.encode("utf-8"), DOC_MIME_TYPES["md"], "md"
    elif target_clean == "txt":
        if src_clean == "html":
            soup = BeautifulSoup(text_content, "html.parser")
            plain = soup.get_text()
        else:
            plain = text_content
        return plain.encode("utf-8"), DOC_MIME_TYPES["txt"], "txt"

    return text_content.encode("utf-8"), DOC_MIME_TYPES.get(target_clean, "text/plain"), target_clean


# --- Helper Functions ---

def _convert_iwork_bundle(input_bytes: bytes, src_clean: str, target_clean: str) -> Tuple[bytes, str, str]:
    """Parses Apple iWork zip archive payload (Pages, Numbers, Keynote)."""
    try:
        with zipfile.ZipFile(io.BytesIO(input_bytes), "r") as zf:
            text_parts = []
            for name in zf.namelist():
                if name.endswith(".xml") or name.endswith(".txt"):
                    content = zf.read(name).decode("utf-8", errors="ignore")
                    text_parts.append(content)
            full_text = "\n\n".join(text_parts) if text_parts else f"Apple {src_clean.upper()} Document Content"
    except Exception:
        full_text = f"Apple {src_clean.upper()} Document Content"

    if target_clean == "pdf":
        return _render_text_to_pdf(full_text), DOC_MIME_TYPES["pdf"], "pdf"
    elif target_clean == "docx":
        return _render_text_to_docx(full_text), DOC_MIME_TYPES["docx"], "docx"
    else:
        return full_text.encode("utf-8"), DOC_MIME_TYPES["txt"], "txt"


def _convert_webloc(input_bytes: bytes, target_clean: str) -> Tuple[bytes, str, str]:
    """Parses Apple Safari .webloc link file."""
    try:
        root = ET.fromstring(input_bytes)
        url = ""
        for elem in root.iter():
            if elem.text and elem.text.startswith("http"):
                url = elem.text.strip()
                break
        if not url:
            url = input_bytes.decode("utf-8", errors="ignore")
    except Exception:
        url = input_bytes.decode("utf-8", errors="ignore")

    return url.encode("utf-8"), DOC_MIME_TYPES["txt"], "txt"


def _extract_text_from_bytes(input_bytes: bytes, src_ext: str) -> str:
    """Safely decodes bytes into text string."""
    try:
        return input_bytes.decode("utf-8")
    except UnicodeDecodeError:
        return input_bytes.decode("latin-1", errors="ignore")


def _convert_from_pdf(input_bytes: bytes, target_clean: str) -> Tuple[bytes, str, str]:
    """Extracts text from PDF and converts to target format."""
    reader = PdfReader(io.BytesIO(input_bytes))
    extracted_text = []
    for page in reader.pages:
        txt = page.extract_text()
        if txt:
            extracted_text.append(txt)
    
    full_text = "\n\n".join(extracted_text) if extracted_text else "No text could be extracted from PDF."

    if target_clean == "docx":
        return _render_text_to_docx(full_text), DOC_MIME_TYPES["docx"], "docx"
    elif target_clean == "html":
        lines = [f"<p>{line}</p>" for line in full_text.split("\n") if line.strip()]
        html = f"<!DOCTYPE html><html><body>{''.join(lines)}</body></html>"
        return html.encode("utf-8"), DOC_MIME_TYPES["html"], "html"
    elif target_clean == "md":
        return full_text.encode("utf-8"), DOC_MIME_TYPES["md"], "md"
    else:
        return full_text.encode("utf-8"), DOC_MIME_TYPES["txt"], "txt"


def _convert_from_docx(input_bytes: bytes, target_clean: str) -> Tuple[bytes, str, str]:
    """Extracts text from Word DOCX and converts to target format."""
    doc = Document(io.BytesIO(input_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text]
    full_text = "\n\n".join(paragraphs)

    if target_clean == "pdf":
        return _render_text_to_pdf(full_text, "txt"), DOC_MIME_TYPES["pdf"], "pdf"
    elif target_clean == "html":
        html_paras = [f"<p>{p}</p>" for p in paragraphs]
        html = f"<!DOCTYPE html><html><body>{''.join(html_paras)}</body></html>"
        return html.encode("utf-8"), DOC_MIME_TYPES["html"], "html"
    elif target_clean == "md":
        return full_text.encode("utf-8"), DOC_MIME_TYPES["md"], "md"
    else:
        return full_text.encode("utf-8"), DOC_MIME_TYPES["txt"], "txt"


def _convert_spreadsheet(input_bytes: bytes, src_clean: str, target_clean: str) -> Tuple[bytes, str, str]:
    """Converts spreadsheets (XLSX, CSV, TSV) using pandas if available, else pure Python."""
    buffer = io.BytesIO(input_bytes)
    if pd is not None:
        if src_clean in ["xlsx", "xls"]:
            df = pd.read_excel(buffer)
        elif src_clean == "tsv":
            df = pd.read_csv(buffer, sep="\t")
        else:
            df = pd.read_csv(buffer)

        out_buffer = io.BytesIO()
        if target_clean == "xlsx":
            with pd.ExcelWriter(out_buffer, engine="openpyxl") as writer:
                df.to_excel(writer, index=False)
            return out_buffer.getvalue(), DOC_MIME_TYPES["xlsx"], "xlsx"
        elif target_clean == "tsv":
            df.to_csv(out_buffer, sep="\t", index=False)
            return out_buffer.getvalue(), DOC_MIME_TYPES["tsv"], "tsv"
        elif target_clean == "html":
            html_str = df.to_html(index=False, classes="table table-striped")
            return html_str.encode("utf-8"), DOC_MIME_TYPES["html"], "html"
        elif target_clean == "md":
            md_str = df.to_markdown(index=False) if hasattr(df, "to_markdown") else df.to_string(index=False)
            return md_str.encode("utf-8"), DOC_MIME_TYPES["md"], "md"
        elif target_clean == "txt":
            txt_str = df.to_string(index=False)
            return txt_str.encode("utf-8"), DOC_MIME_TYPES["txt"], "txt"
        else:
            df.to_csv(out_buffer, index=False)
            return out_buffer.getvalue(), DOC_MIME_TYPES["csv"], "csv"
    else:
        # Pure Python fallback for CSV / TSV without pandas
        import csv
        raw_text = input_bytes.decode("utf-8", errors="ignore")
        src_sep = "\t" if src_clean == "tsv" else ","
        target_sep = "\t" if target_clean == "tsv" else ","
        reader = list(csv.reader(io.StringIO(raw_text), delimiter=src_sep))
        out_buf = io.StringIO()
        writer = csv.writer(out_buf, delimiter=target_sep)
        writer.writerows(reader)
        out_bytes = out_buf.getvalue().encode("utf-8")
        mime = DOC_MIME_TYPES.get(target_clean, "text/csv")
        return out_bytes, mime, target_clean


def _convert_ipynb(input_bytes: bytes, target_clean: str) -> Tuple[bytes, str, str]:
    """Converts Jupyter Notebook (.ipynb) to PY, MD, HTML, TXT, or PDF."""
    try:
        data = json.loads(input_bytes.decode("utf-8"))
    except Exception as e:
        raise ValueError(f"Invalid .ipynb notebook JSON payload: {e}")

    cells = data.get("cells", [])
    py_lines = []
    md_lines = []
    html_parts = []

    for cell in cells:
        cell_type = cell.get("cell_type", "")
        source = "".join(cell.get("source", []))

        if cell_type == "code":
            py_lines.append(source)
            md_lines.append(f"```python\n{source}\n```")
            html_parts.append(f"<pre style='background:#f4f4f4;padding:10px;'><code>{source}</code></pre>")
        elif cell_type == "markdown":
            commented = "\n".join(f"# {line}" for line in source.split("\n"))
            py_lines.append(commented)
            md_lines.append(source)
            html_parts.append(f"<div>{markdown.markdown(source)}</div>")

    full_py = "\n\n# --- Cell ---\n\n".join(py_lines)
    full_md = "\n\n".join(md_lines)
    full_html = f"<!DOCTYPE html><html><body>{''.join(html_parts)}</body></html>"

    if target_clean == "py":
        return full_py.encode("utf-8"), DOC_MIME_TYPES["py"], "py"
    elif target_clean == "html":
        return full_html.encode("utf-8"), DOC_MIME_TYPES["html"], "html"
    elif target_clean == "pdf":
        return _render_text_to_pdf(full_md, "md"), DOC_MIME_TYPES["pdf"], "pdf"
    elif target_clean == "docx":
        return _render_text_to_docx(full_md), DOC_MIME_TYPES["docx"], "docx"
    elif target_clean == "txt":
        return full_py.encode("utf-8"), DOC_MIME_TYPES["txt"], "txt"
    else:
        return full_md.encode("utf-8"), DOC_MIME_TYPES["md"], "md"


def _render_text_to_pdf(text: str, src_format: str = "txt") -> bytes:
    """Renders text string to PDF document using ReportLab."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    normal_style = styles["Normal"]

    story = []
    lines = text.split("\n")
    for line in lines:
        cleaned_line = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").strip()
        if cleaned_line:
            p = Paragraph(cleaned_line, normal_style)
            story.append(p)
        else:
            story.append(Spacer(1, 10))

    if not story:
        story.append(Paragraph("Empty Document", normal_style))

    doc.build(story)
    return buffer.getvalue()


def _render_text_to_docx(text: str) -> bytes:
    """Renders text string to DOCX document using python-docx."""
    doc = Document()
    lines = text.split("\n")
    for line in lines:
        if line.strip():
            doc.add_paragraph(line)
    
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
